#!/usr/bin/env python3
"""Step B0: write docs/tree-review/tree-review.xlsx from docs/tree-review/tree-review.json.

Run after `node scripts/tree-review.js`. The workbook is for review only; the app is unchanged.
Your-call columns (Accept / Override / Ask) come back as data/tree-overrides.json entries.
"""
import json
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..')
SRC = os.path.join(ROOT, 'docs', 'tree-review', 'tree-review.json')
OUT = os.path.join(ROOT, 'docs', 'tree-review', 'tree-review.xlsx')
UNOFFICIAL = 'Unofficial prototype · public data only · not a Region of Peel product'
SCREEN = 'General screening results only. Peel’s W&WW Infrastructure Planning team must validate or confirm all results.'

HEAD = Font(bold=True)
FILL = PatternFill('solid', fgColor='DDE7F3')
CALL = PatternFill('solid', fgColor='FFF2CC')
WRAP = Alignment(wrap_text=True, vertical='top')


def j(v):
    if isinstance(v, list):
        return '; '.join(str(j(x)) for x in v if x not in (None, ''))
    if isinstance(v, dict):
        return ', '.join(f'{k}: {j(x)}' for k, x in v.items())
    if isinstance(v, bool):
        return 'yes' if v else ''
    return v


def table(ws, headers, rows, widths=None, call_cols=0):
    ws.append(headers)
    for c in ws[ws.max_row]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append([j(v) for v in r])
    ws.freeze_panes = ws.cell(row=ws.max_row - len(rows) + 1, column=1) if rows else None
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(48, max(10, len(h) + 2)))
    if call_cols:
        first = len(headers) - call_cols + 1
        start = ws.max_row - len(rows) + 1
        for row in ws.iter_rows(min_row=start, max_row=ws.max_row, min_col=first, max_col=len(headers)):
            for c in row:
                c.fill = CALL
    ws.auto_filter.ref = ws.dimensions


def dict_sheet(wb, title, items, note):
    ws = wb.create_sheet(title[:31])
    ws.append([note]); ws['A1'].font = Font(italic=True)
    if not items:
        ws.append(['(none)']); return
    if isinstance(items, dict):
        items = [{'key': k, **(v if isinstance(v, dict) else {'value': v})} for k, v in items.items()]
    keys = []
    for it in items:
        for k in (it if isinstance(it, dict) else {'value': it}):
            if k not in keys:
                keys.append(k)
    table(ws, keys, [[(it if isinstance(it, dict) else {'value': it}).get(k, '') for k in keys] for it in items])


def main():
    d = json.load(open(SRC))
    wb = Workbook()

    # 1. Read me
    ws = wb.active; ws.title = 'Read me'
    T = d['totals']['Peel']
    lines = [
        ['Tree review (Step B0): applying the containment tree to the unit counts'],
        [UNOFFICIAL], [SCREEN], [],
        [f"Records: data/snapshot.json of {d['snapshotAt'][:10]} · built {d['generatedAt'][:10]} · report only, the app is unchanged."],
        [f"Peel: app {T['appUnits']:,} units → tree {T['treeUnits']:,} ({T['treeUnits'] - T['appUnits']:+,}) · after your overrides {T['afterOverrides']:,}."],
        [f"{T['groups']:,} groups: {T['up']} up, {T['down']} down, {T['unchanged']:,} unchanged. Pre-consultation estimates: {T['preconEstimates']} groups, {T['preconUnits']:,} units (proposed only, never committed)."],
        [],
        ['How to review'],
        ['0. To review: the groups left after your bulk accepts. Only these need a call.'],
        ['1. Changes: every group whose count or grouping changes, largest first, with the call each rule made.'],
        ['2. In the yellow columns choose Accept, Override or Ask. For Override give the units and a reason or source.'],
        ['3. Pre-consultation estimates: counts that come only from a pre-consultation (decided: count them, flagged, never committed). Override any that are not real.'],
        ['4. The other sheets are the audit lists (double counts, mismatches, unlinked nesting, unclear same-land pairs, orphans, low-confidence links, straddling, reversed dates, counts to find).'],
        ['5. Send the workbook back (or list the rows); the decisions go into data/tree-overrides.json and the totals are rebuilt.'],
        [],
        ['Definitions'],
        ['App units: the app’s development units today (the header’s definition: the larger of planned and permitted).'],
        ['Tree units: one count per tree (js/tree.js countTree): each file counts once; children replace a parent that states no count or that they cover ≥90%; same land counts once; condos and pre-consultations count only when nothing else states a count.'],
        ['Group: app development(s) and tree(s) that share files. Kind says whether the tree splits, merges or recounts them.'],
        ['Confidence: the weakest link in the tree (high / medium / low; single file = no links).'],
        ['Plan: the tree’s primary planning area (largest share of the root file’s boundary). Plant: the wastewater plant of the block holding the site.'],
    ]
    for l in lines:
        ws.append(l)
    for r in (1, 9, 15):
        ws.cell(row=r, column=1).font = Font(bold=True, size=13 if r == 1 else 11)
    ws.column_dimensions['A'].width = 150

    # 2. Summary
    ws = wb.create_sheet('Summary')
    rows = []
    for m in ['Mississauga', 'Brampton', 'Caledon', 'Peel']:
        t = d['totals'].get(m)
        if not t:
            continue
        ch = t['treeUnits'] - t['appUnits']
        rows.append([m, t['groups'], t['appUnits'], t['treeUnits'], ch, round(ch / t['appUnits'] * 100, 1) if t['appUnits'] else '', t['afterOverrides'], t['up'], t['down'], t['unchanged'], t['preconEstimates'], t['preconUnits']])
    table(ws, ['Municipality', 'Groups', 'App units', 'Tree units', 'Change', 'Change %', 'After overrides', 'Up', 'Down', 'Unchanged', 'Pre-con estimates', 'Pre-con units'], rows)
    ws.append([]); ws.append(['By kind of change']); ws.cell(row=ws.max_row, column=1).font = HEAD
    ws.append(['Kind', 'Groups', 'Net change (units)'])
    for c in ws[ws.max_row]:
        c.font, c.fill = HEAD, FILL
    for k, v in sorted(d['byKind'].items(), key=lambda kv: -abs(kv[1]['change'])):
        ws.append([k, v['groups'], v['change']])

    # 3. Changes (with your-call columns)
    heads = ['Group', 'Municipality', 'Kind', 'App development(s)', 'App stage', 'Committed now (D-5-1)', 'App units', 'Tree units', 'Change', 'Flags',
             'Tree root file(s)', 'Counting rule', 'Confidence', 'Plan (primary)', 'Plan share %', 'Plant', 'Block', 'Files in tree(s)',
             'Your call', 'Why this call', 'Override units', 'Reason / source']
    widths = {'Why this call': 44, 'App development(s)': 34, 'Counting rule': 46, 'Files in tree(s)': 40, 'Flags': 30, 'Kind': 30, 'Plan (primary)': 28, 'Reason / source': 40, 'Tree root file(s)': 24}

    def row_of(r):
        ts = r['trees']
        return [r['group'], r['municipality'], r['kind'], r['appDevelopments'], r['appPhase'], r['committedNow'], r['appUnits'], r['treeUnits'], r['change'], r['flags'],
                [t['root'] for t in ts], [f"{t['root']}: {t['units']:,} ({t['basis']})" for t in ts], [t['confidence'] for t in ts],
                [t['plan'] for t in ts], [t['planShare'] for t in ts if t['planShare'] is not None], [t['plant'] for t in ts], [t['block'] for t in ts],
                [', '.join(t['files'][:12]) + (f' +{len(t["files"]) - 12}' if len(t['files']) > 12 else '') for t in ts],
                r.get('call', ''), r.get('callWhy', ''), next((t['override'] for t in ts if t['override'] is not None), ''), next((t['overrideReason'] for t in ts if t['overrideReason']), '')]

    ask = [r for r in d['rows'] if r.get('call') == 'Ask']
    ws = wb.create_sheet('To review', 1)
    ws.append([f'{len(ask)} groups left for you after the bulk accepts of 2026-10-15 (data/tree-overrides.json "accepted"). Change Ask to Accept or Override; for Override give units and a reason.'])
    ws['A1'].font = Font(italic=True)
    table(ws, heads, [row_of(r) for r in ask], widths, call_cols=4)
    dv0 = DataValidation(type='list', formula1='"Accept,Override,Ask"', allow_blank=True)
    ws.add_data_validation(dv0); dv0.add(f'S3:S{ws.max_row}')
    changed = [r for r in d['rows'] if r['kind'] != 'unchanged']
    ws = wb.create_sheet('Changes')
    table(ws, heads, [row_of(r) for r in changed], widths, call_cols=4)
    dv = DataValidation(type='list', formula1='"Accept,Override,Ask"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(f'S2:S{ws.max_row}')

    # 4. Pre-consultation estimates
    pre = [r for r in d['rows'] if any(t['preconOnly'] for t in r['trees'])]
    ws = wb.create_sheet('Pre-consultation estimates')
    table(ws, heads, [row_of(r) for r in pre], widths, call_cols=4)
    dv2 = DataValidation(type='list', formula1='"Accept,Override,Ask"', allow_blank=True)
    ws.add_data_validation(dv2); dv2.add(f'S2:S{ws.max_row}')

    # 5. Every group (unchanged included), for lookups
    ws = wb.create_sheet('All groups')
    table(ws, heads[:18], [row_of(r)[:18] for r in d['rows']], widths)

    # 6+. Audit lists
    L = d['lists']
    dict_sheet(wb, 'Double counts', (L['doubleCounts'].get('currentFileInTwoDevelopments') or []) + [{'file': x} if not isinstance(x, dict) else x for x in (L['doubleCounts'].get('proposedFileInTwoTrees') or [])], 'Files the app puts in two developments, and files claimed by two trees.')
    dict_sheet(wb, 'Unit mismatches', L['unitMismatches'], 'Children counting more than their parent by over 10%, or counted files not linked directly.')
    dict_sheet(wb, 'Nested not linked', L['nestedFormalNotLinked'], 'A formal file inside another formal file with no link between them: should the inner one be a child?')
    dict_sheet(wb, 'Same-land unclear', L['sameLandUnclear'], 'Same-land site plans the scoring could not call (phase / resubmission / redevelopment): both are counted until you decide.')
    dict_sheet(wb, 'Condo-only trees', L['condoOnlyTrees'], 'Trees whose only count is a condominium (no site plan or draft plan states one).')
    dict_sheet(wb, 'Orphans 50+', L['orphans'], 'Formal files with 50+ units that link to nothing.')
    dict_sheet(wb, 'Low-confidence links', L['lowConfidenceLinks'], 'Links made with low confidence (small overlap, address mismatch or date order).')
    dict_sheet(wb, 'Straddling plans', L['straddling'], 'Trees with no plan holding 90% of the root file.')
    dict_sheet(wb, 'Dates reversed', L['dateReversed'], 'A child file dated before its parent.')
    dict_sheet(wb, 'Counts to find', L['growthToFind'], 'Active developments where no file states a count and the wording does not settle it.')
    wb.save(OUT)
    print(f'wrote {OUT}: {len(changed)} changes, {len(pre)} pre-consultation estimates, {len(d["rows"])} groups')


if __name__ == '__main__':
    main()
